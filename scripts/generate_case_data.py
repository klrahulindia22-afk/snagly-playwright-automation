"""Generate the committed 172-case data registry from the approved XLSX catalogue."""

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

ROLE_OVERRIDES = {
    "SNAG-TC-045": "team",
    "SNAG-TC-106": "client",
    "SNAG-TC-117": "client",
    "SNAG-TC-167": "client",
}


def infer_role(preconditions: str, scenario: str) -> str:
    text = f"{preconditions} {scenario}".lower()
    if "logged in as admin" in text or "admin account" in text:
        return "admin"
    if "logged in as client" in text or "client user performs" in text:
        return "client"
    if "logged in as team" in text or "team user performs" in text:
        return "team"
    if "logged in as other owner" in text:
        return "other_owner"
    return "owner"


def infer_refs(module: str, test_type: str, scenario: str, test_data: str) -> list[str]:
    text = f"{module} {test_type} {scenario} {test_data}".lower()
    refs = []
    mapping = {
        "authentication": "emails",
        "password": "passwords",
        "email": "emails",
        "board": "boards",
        "list": "lists",
        "card": "cards",
        "label": "labels",
        "attachment": "attachments",
        "file": "attachments",
        "due date": "dates",
        "timezone": "dates",
        "responsive": "viewports",
        "viewport": "viewports",
        "report": "report_periods",
        "integration": "integration_errors",
        "security": "security_payloads",
        "injection": "security_payloads",
        "xss": "security_payloads",
        "boundary": "text_boundaries",
        "validation": "text_boundaries",
    }
    for needle, ref in mapping.items():
        if needle in text and ref not in refs:
            refs.append(ref)
    return refs


def generate(source: Path, output: Path) -> None:
    sheet = load_workbook(source, read_only=True, data_only=True)["Test Cases"]
    rows = sheet.iter_rows()
    headers = [cell.value for cell in next(rows)]
    index = {name: position for position, name in enumerate(headers)}
    cases = []
    for cells in rows:
        row = tuple(cell.value for cell in cells)
        if not row[index["Test Case ID"]]:
            continue
        module = str(row[index["Module"]] or "")
        test_type = str(row[index["Type"]] or "")
        scenario = str(row[index["Test Scenario"]] or "")
        preconditions = str(row[index["Preconditions"]] or "")
        case_id = str(row[index["Test Case ID"]])
        source_test_data = str(row[index["Test Data"]] or "")
        cases.append({
            "case_id": case_id,
            "module": module,
            "submodule": row[index["Submodule"]],
            "type": test_type,
            "priority": row[index["Priority"]],
            "scenario": scenario,
            "preconditions": preconditions,
            "source_test_data": source_test_data,
            "role": ROLE_OVERRIDES.get(case_id, infer_role(preconditions, scenario)),
            "review_status": "generated_needs_review",
            "data_refs": infer_refs(module, test_type, scenario, source_test_data),
            "values": {},
        })
    payload = {"schema_version": 1, "expected_count": len(cases), "cases": cases}
    output.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Generated {len(cases)} case records in {output}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path, default=Path("test_data/cases.json"))
    args = parser.parse_args()
    generate(args.source, args.output)
